import pandas as pd
from docx import Document
from pathlib import Path

# 1. Criar a Planilha de Dados
dados = {
    'Cliente': ['Tech Solutions', 'Mercado Do Bairro', 'Advocacia Silva', 'Startup Inova'],
    'CNPJ': ['12.345.678/0001-90', '98.765.432/0001-10', '11.222.333/0001-44', '55.666.777/0001-88'],
    'Valor': [15000.00, 2500.50, 8000.00, 50000.00],
    'Prazo': ['12 meses', '6 meses', '24 meses', 'Indeterminado'],
    'Vendedor': ['Ana', 'Carlos', 'Ana', 'Carlos'] # Usaremos isso para separar as pastas
}

df = pd.read_excel('contratos.xlsx') if Path('contratos.xlsx').exists() else pd.DataFrame(dados)
df.to_excel("contratos.xlsx", index=False)

# 2. Criar a Minuta Padrão (Template do Word)
document = Document()
document.add_heading('CONTRATO DE PRESTAÇÃO DE SERVIÇOS', 0)

p = document.add_paragraph('Pelo presente instrumento, a empresa ')
p.add_run('{{CLIENTE}}').bold = True
p.add_run(', inscrita no CNPJ sob nº ')
p.add_run('{{CNPJ}}').bold = True
p.add_run(', doravante denominada CONTRATANTE.')

document.add_paragraph('CLÁUSULA 1: O valor deste contrato é de {{VALOR}}, a ser pago mensalmente.')
document.add_paragraph('CLÁUSULA 2: A vigência deste contrato será de {{PRAZO}}, iniciando-se na data de assinatura.')
document.add_paragraph('\n\n___________________________\nAssinatura do Responsável')

document.save('Minuta_Padrao.docx')

print("✅ Ambiente preparado: 'contratos.xlsx' e 'Minuta_Padrao.docx' criados.")